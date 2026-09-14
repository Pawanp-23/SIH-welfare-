# -*- coding: utf-8 -*-
"""Build the SIH 2026 idea PPT for PS SIH26186 on top of the official template."""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from pptx.text.text import _Run
from lxml import etree

TEMPLATE = "template.pptx"
OUT = "SAHARA_SIH26186_Idea_PPT.pptx"

# ---- palette -------------------------------------------------------------
NAVY = "1B2A4A"; TEAL = "0E7C7B"; SAFF = "E8772E"
LIGHT = "F3F5F8"; TEXT = "1F2933"; MUTED = "5B6573"; WHITE = "FFFFFF"
TEAL_LT = "E3F2F1"; SAFF_LT = "FDEBDD"; NAVY_LT = "E6EAF2"; LINE = "D5DAE1"
NAVY_MID = "2E4270"
FONT = "Calibri"

TEAM_NAME = "Team Name"      # <-- fill in
TEAM_ID = "________"         # <-- fill in

def rgb(h): return RGBColor.from_string(h)

# ---- helpers -------------------------------------------------------------
def rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75, shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s

def rrect(slide, x, y, w, h, fill=None, line=None, radius=0.08, **kw):
    return rect(slide, x, y, w, h, fill, line, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius, **kw)

def oval(slide, x, y, d, fill, line=None):
    return rect(slide, x, y, d, d, fill, line, shape=MSO_SHAPE.OVAL)

def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, margin=0.03, wrap=True):
    """paras: list of dicts {t|runs, size, bold, color, italic, align, bullet, after, before, line, font}"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = anchor
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        if p.get("after") is not None: para.space_after = Pt(p["after"])
        if p.get("before") is not None: para.space_before = Pt(p["before"])
        if p.get("line") is not None: para.line_spacing = p["line"]
        runs = p.get("runs") or [dict(t=p.get("t", ""))]
        for r in runs:
            run = para.add_run()
            run.text = r["t"]
            f = run.font
            f.name = r.get("font", p.get("font", FONT))
            f.size = Pt(r.get("size", p.get("size", 11)))
            f.bold = r.get("bold", p.get("bold", False))
            f.italic = r.get("italic", p.get("italic", False))
            f.color.rgb = rgb(r.get("color", p.get("color", TEXT)))
            if r.get("url"):
                run.hyperlink.address = r["url"]
        if p.get("bullet") or p.get("indent"):
            pPr = para._p.get_or_add_pPr()
            ind = p.get("indent", 0.16)
            pPr.set("marL", str(int(Inches(ind))))
            if p.get("bullet"):
                pPr.set("indent", str(-int(Inches(ind))))
                if p.get("bucolor"):
                    clr = etree.SubElement(pPr, qn("a:buClr")); srgb = etree.SubElement(clr, qn("a:srgbClr")); srgb.set("val", p["bucolor"])
                bf = etree.SubElement(pPr, qn("a:buFont")); bf.set("typeface", "Arial")
                bc = etree.SubElement(pPr, qn("a:buChar")); bc.set("char", p.get("char", "•"))
    return tb

def icon(slide, name, x, y, size):
    return slide.shapes.add_picture(f"icons/{name}.png", Inches(x), Inches(y), Inches(size), Inches(size))

def icon_circle(slide, name, cx, cy, d, fill, icon_ratio=0.52):
    oval(slide, cx - d/2, cy - d/2, d, fill)
    s = d * icon_ratio
    icon(slide, name, cx - s/2, cy - s/2, s)

def arrow(slide, x, y, w, h, fill, direction="right"):
    shp = MSO_SHAPE.RIGHT_ARROW if direction == "right" else MSO_SHAPE.DOWN_ARROW
    return rect(slide, x, y, w, h, fill, shape=shp)

def line(slide, x1, y1, x2, y2, color, w=1.0, dash=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(w)
    if dash: c.line.dash_style = dash
    return c

def set_title(slide, txt):
    for sh in slide.shapes:
        if sh.name == "Title 1":
            sh.text_frame.paragraphs[0].runs[-1].text = txt
            return
    raise RuntimeError("no title")

def set_oval(slide, txt):
    for sh in slide.shapes:
        if sh.name.startswith("Oval"):
            r = sh.text_frame.paragraphs[0].runs[0]; r.text = txt; r.font.size = Pt(12); r.font.bold = True
            return

def drop_shape(slide, name):
    for sh in list(slide.shapes):
        if sh.name == name:
            sh._element.getparent().remove(sh._element)

def delete_slide(prs, idx):
    sldIdLst = prs.slides._sldIdLst
    sld = list(sldIdLst)[idx]
    prs.part.drop_rel(sld.rId)
    sldIdLst.remove(sld)

# ==========================================================================
prs = Presentation(TEMPLATE)
delete_slide(prs, 6)                       # remove "Important instructions"
S = prs.slides

# ---------------------------------------------------------------- SLIDE 1
s = S[0]
fields = [
    ("Problem Statement ID – ", "SIH26186", 22),
    ("Problem Statement Title – ", "AI-Based Predictive Personnel Stress and Welfare Monitoring System for Uniformed Forces", 17),
    ("Theme – ", "MedTech / BioTech / HealthTech", 22),
    ("PS Category – ", "Software", 22),
    ("Team ID – ", TEAM_ID, 22),
    ("Team Name – ", TEAM_NAME, 22),
]
for sh in s.shapes:
    if sh.name == "TextBox 9":
        tf = sh.text_frame
        paras = tf.paragraphs[1:]           # paragraph 0 is an empty spacer
        for para, (label, val, sz) in zip(paras, fields):
            r0 = para.runs[0]
            r0.text = label; r0.font.size = Pt(sz); r0.font.bold = True
            r1 = copy.deepcopy(r0._r); r0._r.addnext(r1)
            para._p.pPr.set('algn', 'l')
            run = _Run(r1, para); run.text = val; run.font.bold = False; run.font.size = Pt(sz)
            run.font.color.rgb = rgb(NAVY)
            lnSpc = para._p.pPr.find(qn("a:lnSpc"))
            if lnSpc is not None: lnSpc.find(qn("a:spcPct")).set("val", "140000")
        sh.top = Inches(2.05)

# ---------------------------------------------------------------- SLIDE 2
s = S[1]
set_title(s, "SAHARA – AI Welfare Intelligence"); set_oval(s, TEAM_NAME); drop_shape(s, "TextBox 8")
LX, LW = 0.4, 7.55
text(s, LX, 1.22, LW, 0.35, [dict(t="Predictive Stress & Welfare Monitoring for Uniformed Forces  •  welfare-first  •  privacy-first  •  explainable",
                                  size=12, italic=True, color=TEAL, bold=True)])

def section(slide, y, iconname, title, body_paras, body_h, gap=0.05):
    icon_circle(slide, iconname, LX + 0.2, y + 0.2, 0.4, TEAL_LT, 0.5)
    text(slide, LX + 0.5, y, LW - 0.5, 0.4, [dict(t=title, size=14, bold=True, color=NAVY)], anchor=MSO_ANCHOR.MIDDLE)
    text(slide, LX + 0.05, y + 0.42 + gap, LW - 0.1, body_h, body_paras)

section(s, 1.62, "target_t", "Proposed Solution",
    [dict(t="SAHARA (सहारा = support) is a privacy-first AI platform that fuses HRMS context (duty hours, leave, deployment, night shifts), "
            "a 60-second daily check-in and opt-in wearable signals (sleep, HRV) into a 0–100 Welfare Risk Index with a 7-day forecast — "
            "explains every score with exact SHAP attributions and puts a specific, evidenced action before the officer who can authorise it. "
            "Welfare actions, never disciplinary flags.", size=11, line=1.05)], 0.95)

section(s, 2.98, "lightbulb_t", "How It Addresses the Problem", [
    dict(t="Replaces manual observation with continuous early warning — predicts deterioration 7 days ahead with a calibrated escalation probability, before an incident, not after", size=11, bullet=True, bucolor=TEAL, after=3),
    dict(t="Welfare Officer gets ranked cases with the “why”; Command sees only pseudonymised unit aggregates — an individual record returns 403 and is audited", size=11, bullet=True, bucolor=TEAL, after=3),
    dict(t="Recommends interventions ranked by measured counterfactual effect (rest day, break night-shift run, sanction leave, peer support) with SOP citation + what-if simulator", size=11, bullet=True, bucolor=TEAL),
], 1.15)

section(s, 4.74, "rocket_t", "Innovation & Uniqueness", [
    dict(t="Consent-gated, tiered data (check-in → HRMS → wearables) with an append-only consent ledger (DPDP Act 2023) — every jawan controls what is shared", size=11, bullet=True, bucolor=SAFF, after=3),
    dict(t="Exact TreeSHAP on every prediction (~320 µs) + 80 % prediction interval; escalation by 4 inspectable rules — AI ranks, the welfare officer decides", size=11, bullet=True, bucolor=SAFF, after=3),
    dict(t="Stigma-safe by design: HMAC pseudonyms, k = 10 suppression, welfare-only RBAC, hash-chained audit — never linked to ACR / promotion / posting", size=11, bullet=True, bucolor=SAFF, after=3),
    dict(t="Model runs in TypeScript (Python parity verified to 1e-13); zero native dependencies — runs on any laptop, on-prem ready", size=11, bullet=True, bucolor=SAFF),
], 1.55)

# right "at a glance" card
CX, CY, CW, CH = 8.25, 1.3, 4.7, 5.5
rrect(s, CX, CY, CW, CH, NAVY, radius=0.05)
text(s, CX, CY + 0.12, CW, 0.35, [dict(t="SAHARA AT A GLANCE", size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)])
text(s, CX, CY + 0.5, CW, 0.25, [dict(t="SIGNALS  (consent-tiered, opt-in)", size=8.5, color="A9B7D0", align=PP_ALIGN.CENTER)])
inputs = [("db", "HRMS Context", "duty hrs · leave · deployment"), ("mobile", "Daily Check-in", "5 questions · 60 seconds"), ("watch", "Wearables (opt-in)", "sleep · HRV · activity")]
bw, bg = 1.42, 0.12
bx = CX + (CW - (3*bw + 2*bg)) / 2
for i, (ic, t1, t2) in enumerate(inputs):
    x = bx + i*(bw+bg); y = CY + 0.8
    rrect(s, x, y, bw, 1.0, NAVY_MID, radius=0.1)
    icon(s, ic, x + bw/2 - 0.16, y + 0.1, 0.32)
    text(s, x, y + 0.45, bw, 0.55, [dict(t=t1, size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER), dict(t=t2, size=7.5, color="C9D3E6", align=PP_ALIGN.CENTER)])
arrow(s, CX + CW/2 - 0.14, CY + 1.87, 0.28, 0.3, TEAL, "down")
ey = CY + 2.22
rrect(s, CX + 0.28, ey, CW - 0.56, 1.0, TEAL, radius=0.1)
icon(s, "brain", CX + 0.45, ey + 0.28, 0.45)
text(s, CX + 1.0, ey + 0.05, CW - 1.3, 0.9, [dict(t="Predictive Risk Engine", size=10.5, bold=True, color=WHITE),
     dict(t="Welfare Risk Index 0–100 · 7-day forecast · P(escalation) · exact SHAP drivers · ranked interventions", size=8.5, color="E3F2F1")], anchor=MSO_ANCHOR.MIDDLE)
arrow(s, CX + CW/2 - 0.14, ey + 1.07, 0.28, 0.3, TEAL, "down")
oy = ey + 1.42
outs = [("usershield", "Welfare Officer", "ranked cases · evidenced actions · what-if simulator"), ("chart", "Command / HQ", "unit aggregates & escalation velocity — pseudonyms only")]
ow = (CW - 0.56 - 0.15) / 2
for i, (ic, t1, t2) in enumerate(outs):
    x = CX + 0.28 + i*(ow + 0.15)
    rrect(s, x, oy, ow, 1.0, NAVY_MID, radius=0.1)
    icon(s, ic, x + 0.12, oy + 0.33, 0.34)
    text(s, x + 0.5, oy + 0.04, ow - 0.55, 0.95, [dict(t=t1, size=9.5, bold=True, color=WHITE), dict(t=t2, size=8, color="C9D3E6")], anchor=MSO_ANCHOR.MIDDLE)
py = oy + 1.12
rrect(s, CX + 0.28, py, CW - 0.56, 0.72, None, SAFF, radius=0.1, line_w=1.25)
icon(s, "shield", CX + 0.42, py + 0.19, 0.34)
text(s, CX + 0.85, py + 0.02, CW - 1.15, 0.68, [dict(t="Privacy Shield", size=9.5, bold=True, color=SAFF),
     dict(t="consent ledger · HMAC pseudonyms · k = 10 suppression · RBAC · hash-chained audit · DPDP Act 2023", size=8, color="E3F2F1")], anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- SLIDE 3
import math
s = S[2]
set_title(s, "TECHNICAL APPROACH"); set_oval(s, TEAM_NAME); drop_shape(s, "TextBox 8")
text(s, 2.0, 1.1, 9.33, 0.3, [dict(t="System architecture & technology stack", size=13, italic=True, color=MUTED, align=PP_ALIGN.CENTER)])

SPOKE = ["2E6FD9", "E8772E", "0E7C7B", "3DA35D", "7B4FBF", "D9463E"]
HX, HY, RD = 4.9, 4.05, 3.0
spokes = [
    ("sp_ui",     "Frontend & Role-gated UI",   "React 19 + Vite dashboard — 9 screens: daily check-in, my welfare, casework, force view, what-if simulator, model card"),
    ("sp_api",    "Secure API & Identity",       "Express 4 on Node 22; scrypt + HMAC-JWT; RBAC for personnel / welfare officer / command / admin; per-route rate limits"),
    ("sp_ml",     "ML Risk Engine",              "5 gradient-boosted ensembles on 25 features → Welfare Risk Index 0–100, 7-day forecast, P(escalation in 14 d)"),
    ("sp_db",     "Data & Consent Layer",        "node:sqlite, 11 tables; append-only consent ledger (DPDP); HRMS import + 5-question daily check-in + opt-in wearables"),
    ("sp_shap",   "Explainability & Escalation", "Exact TreeSHAP on every score (~320 µs); 80 % prediction interval; 4 documented rules open a case — not a black box"),
    ("sp_shield", "Privacy, Audit & Alerts",     "HMAC pseudonyms; k = 10 suppression; SHA-256 hash-chained audit; live SSE alerts to the assigned welfare officer"),
]
angles = [330, 270, 210, 30, 90, 150]      # clockwise from 12 o'clock
rows_y = [1.5, 3.1, 4.7]
row_h = 1.35
# connectors first (behind everything)
for i, a in enumerate(angles):
    r = RD/2 - 0.3
    ix, iy = HX + r*math.sin(math.radians(a)), HY - r*math.cos(math.radians(a))
    ry = rows_y[i % 3] + 0.55
    if i < 3:
        line(s, 3.2, ry, ix, iy, SPOKE[i], 1.25)
        oval(s, 3.2 - 0.05, ry - 0.05, 0.1, SPOKE[i])
    else:
        line(s, 6.6, ry, ix, iy, SPOKE[i], 1.25)
        oval(s, 6.6 - 0.05, ry - 0.05, 0.1, SPOKE[i])
# ring + hub
icon(s, "ring", HX - RD/2, HY - RD/2, RD)
text(s, HX - 0.85, HY - 0.4, 1.7, 0.8, [dict(t="SYSTEM\nARCHITECTURE", size=11.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER, line=1.0)], anchor=MSO_ANCHOR.MIDDLE)
for i, a in enumerate(angles):
    r = RD/2 - 0.3
    ix, iy = HX + r*math.sin(math.radians(a)), HY - r*math.cos(math.radians(a))
    oval(s, ix - 0.33, iy - 0.33, 0.66, WHITE)
    icon_circle(s, spokes[i][0], ix, iy, 0.54, SPOKE[i], 0.5)
# numbered items
for i, (ic, t1, t2) in enumerate(spokes):
    y = rows_y[i % 3]
    if i < 3:
        x, w, al = 0.4, 2.75, PP_ALIGN.RIGHT
    else:
        x, w, al = 6.6, 2.85, PP_ALIGN.LEFT
    text(s, x, y, w, row_h, [
        dict(t=str(i+1), size=24, bold=True, color=SPOKE[i], align=al, line=0.9, after=0),
        dict(t=t1, size=11.5, bold=True, color=NAVY, align=al, after=2),
        dict(t=t2, size=8.5, color=TEXT, align=al, line=1.0),
    ])
# verified-model strip
rrect(s, 0.4, 6.08, 9.0, 0.72, LIGHT, radius=0.1)
text(s, 0.5, 6.1, 2.2, 0.68, [dict(t="VERIFIED MODEL", size=9.5, bold=True, color=NAVY, after=1),
                              dict(t="held-out, split by individual · npm run ml:verify", size=7.5, color=MUTED)], anchor=MSO_ANCHOR.MIDDLE)
metrics = [("3.87", "WRI MAE"), ("90.3 %", "band accuracy"), ("0.812", "escalation ROC-AUC"), ("0.78", "recall (F2-tuned)"), ("79.5 %", "80 % PI coverage"), ("≤ 1e-13", "TS ↔ Python parity")]
mx = 2.7; mw = 1.1
for j, (v, l) in enumerate(metrics):
    text(s, mx + j*mw, 6.1, mw, 0.68, [dict(t=v, size=12, bold=True, color=TEAL, align=PP_ALIGN.CENTER, after=0),
                                       dict(t=l, size=7.5, color=MUTED, align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE)
# right panel — core technology flow
PX, PY, PW, PH = 9.65, 1.3, 3.3, 5.5
rrect(s, PX, PY, PW, PH, LIGHT, radius=0.05)
text(s, PX, PY + 0.08, PW, 0.3, [dict(t="SAHARA – Core Technology Flow", size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER)])
techrows = [
    ("ly_front",   "Frontend",       [("b_react", "React 19 + TypeScript 5.8"), ("b_vite", "Vite 6 · Tailwind CSS 4")]),
    ("ly_chart",   "UI & Charts",    [("b_framer", "Motion 12 (reduced-motion aware)"), ("b_chart", "Recharts 3 · CVD-safe palette")]),
    ("ly_back",    "Backend",        [("b_node", "Node 22 · Express 4"), ("b_sse", "SSE live events · rate limits")]),
    ("ly_train",   "ML Training",    [("b_python", "Python · scikit-learn"), ("b_pandas", "pandas · numpy (offline only)")]),
    ("ly_runtime", "ML Runtime",     [("b_ts", "TypeScript GBM inference"), ("b_shap", "Exact TreeSHAP ≈ 320 µs")]),
    ("ly_db",      "Database",       [("b_sqlite", "node:sqlite (WAL) · 11 tables"), ("b_hash", "SHA-256 hash-chained audit")]),
    ("ly_sec",     "Security",       [("b_jwt", "scrypt · HMAC-JWT · RBAC"), ("b_key", "HMAC pseudonyms · k = 10")]),
    ("ly_llm",     "LLM (optional)", [("b_gemini", "Gemini 2.0 Flash over SOP RAG")]),
    ("ly_deploy",  "Build & Deploy", [("b_esbuild", "Vite + esbuild → one Node process"), ("b_docker", "Docker · on-prem / NIC MeghRaj")]),
]
ry = PY + 0.45; rh = 0.55
for k, (lic, lab, techs) in enumerate(techrows):
    if k > 0:
        line(s, PX + 0.12, ry - 0.02, PX + PW - 0.12, ry - 0.02, LINE, 0.5)
    icon(s, lic, PX + 0.12, ry + 0.15, 0.22)
    text(s, PX + 0.38, ry, 1.05, rh, [dict(t=lab, size=8.5, bold=True, color=NAVY, line=1.0)], anchor=MSO_ANCHOR.MIDDLE)
    n = len(techs)
    for j, (bic, t) in enumerate(techs):
        ty = ry + (rh/2 - 0.12) if n == 1 else ry + 0.05 + j*0.25
        icon(s, bic, PX + 1.45, ty + 0.02, 0.18)
        text(s, PX + 1.68, ty - 0.02, PW - 1.75, 0.26, [dict(t=t, size=7.5, color=TEXT)], anchor=MSO_ANCHOR.MIDDLE)
    ry += rh

# ---------------------------------------------------------------- SLIDE 4
s = S[3]
set_title(s, "FEASIBILITY AND VIABILITY"); set_oval(s, TEAM_NAME); drop_shape(s, "TextBox 8")
cols = [
    ("check_w", TEAL, TEAL_LT, "Feasibility", "check_t", [
        "Software-only — no new hardware; wearables strictly optional",
        "Leave, duty & transfer data already digitised in CAPF PIS / e-HRMS",
        "Working prototype: real trained model (5 GBM ensembles, 25 features), 9 role-gated screens, end-to-end verified",
        "Zero native dependencies — one npm run dev; no Python at runtime; SQLite → PostgreSQL swap when scaling",
        "Aligned with MHA Task Force (2022) mandate on stress & suicide prevention",
    ]),
    ("warn_w", SAFF, SAFF_LT, "Challenges & Risks", "excl_o", [
        "Trust & stigma — fear that a risk flag will affect career or posting",
        "Cyber threats on highly sensitive psychological & welfare data",
        "False positives / negatives; cold-start with little labelled data",
        "Low self-reporting in remote, low-connectivity deployments",
        "Cultural & linguistic diversity across forces and units",
    ]),
    ("tools_w", NAVY, NAVY_LT, "Mitigation Strategies", "arrow_n", [
        "Welfare-only charter: no link to ACR/promotion; jawan can view & contest own score",
        "scrypt + HMAC-JWT, RBAC, HMAC pseudonyms, hash-chained audit, on-prem hosting, DPDP-compliant",
        "Human-in-the-loop; threshold tuned on F2 for recall (0.78); model card, calibration & fairness audit in-app",
        "60-second, 5-question check-in; HRMS + wearables carry the signal (self-report-only MAE 10.8 vs 3.9)",
        "Dataset-agnostic pipeline: retrains on a real HRMS CSV without code change; drift (PSI) monitored",
    ]),
]
cw, cg, y0, hh, bh = 2.6, 0.15, 1.32, 0.48, 3.02
for i, (hic, hcol, bcol, title, bic, items) in enumerate(cols):
    x = 0.4 + i*(cw + cg)
    rrect(s, x, y0, cw, hh, hcol, radius=0.15)
    icon(s, hic, x + 0.15, y0 + 0.12, 0.26)
    text(s, x + 0.48, y0, cw - 0.5, hh, [dict(t=title, size=13, bold=True, color=WHITE)], anchor=MSO_ANCHOR.MIDDLE)
    rrect(s, x, y0 + hh + 0.06, cw, bh, bcol, radius=0.05)
    iy = y0 + hh + 0.16
    for it in items:
        icon(s, bic, x + 0.12, iy + 0.04, 0.16)
        text(s, x + 0.36, iy - 0.02, cw - 0.44, 0.62, [dict(t=it, size=9, color=TEXT, line=1.0)])
        iy += 0.58
# viability tiles
VX = 8.8; tw, tg, th, trg = 2.0, 0.15, 1.62, 0.3
tiles = [("rupee_n", "₹0 hardware", "Software-only platform; opt-in wearables use devices personnel already own"),
         ("code_n", "Zero licence cost", "100 % open-source stack; indigenous IP retained by MHA / forces"),
         ("cloud_n2", "Existing infra", "Single Node process; runs on-prem or NIC MeghRaj; air-gap ready"),
         ("handshake_n", "Opt-in & DPDP-ready", "Consent-based tiers — no new legal mandate needed to start a pilot")]
for i, (ic, t1, t2) in enumerate(tiles):
    r, c = divmod(i, 2)
    x = VX + c*(tw + tg); y = y0 + r*(th + trg)
    rrect(s, x, y, tw, th, WHITE, LINE, radius=0.08)
    icon(s, ic, x + 0.15, y + 0.15, 0.36)
    text(s, x + 0.12, y + 0.55, tw - 0.24, 1.05, [dict(t=t1, size=12, bold=True, color=NAVY, after=2), dict(t=t2, size=8.5, color=MUTED, line=1.0)])
# roadmap
text(s, 0.4, 4.98, 6, 0.3, [dict(t="PHASED ROLLOUT ROADMAP", size=12, bold=True, color=NAVY)])
phases = [("Phase 0 · 0–2 mo", "Prototype built on synthetic cohort (2,200 personnel x 12 wk); privacy & security review", NAVY),
          ("Phase 1 · 2–6 mo", "90-day pilot in one battalion; validate lead-time & false-positive rate", TEAL),
          ("Phase 2 · 6–12 mo", "Roll out across one force (e.g., CRPF) with live HRMS integration", TEAL),
          ("Phase 3 · 12–24 mo", "All CAPFs, Armed Forces & State Police; cross-force benchmarking", SAFF)]
pw, py, ph = 3.28, 5.32, 1.42
for i, (t1, t2, col) in enumerate(phases):
    x = 0.4 + i*(pw - 0.19)
    shp = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
    sp = rect(s, x, py, pw, ph, col, shape=shp)
    sp.adjustments[0] = 0.22
    text(s, x + (0.15 if i == 0 else 0.5), py + 0.12, pw - 0.95, ph - 0.24,
         [dict(t=t1, size=10.5, bold=True, color=WHITE, after=2), dict(t=t2, size=8.5, color=WHITE, line=1.0)], anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- SLIDE 5
s = S[4]
set_title(s, "IMPACT AND BENEFITS"); set_oval(s, TEAM_NAME); drop_shape(s, "TextBox 8")
stats = [("730", "suicides in CAPFs, NSG & Assam Rifles in 2020–24 (MHA, Rajya Sabha, Dec 2024)"),
         ("55,500+", "voluntary retirements & resignations in the same period — stress and attrition are linked"),
         ("Reactive → Proactive", "SAHARA moves welfare from post-incident response to early, evidence-based prevention")]
sw, sg, sy, shh = 4.05, 0.2, 1.3, 0.85
for i, (n, l) in enumerate(stats):
    x = 0.4 + i*(sw + sg)
    rrect(s, x, sy, sw, shh, LIGHT, radius=0.1)
    nw = 1.55 if i < 2 else 2.3
    text(s, x + 0.1, sy, nw, shh, [dict(t=n, size=22 if i < 2 else 15, bold=True, color=SAFF, align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + nw + 0.1, sy + 0.05, sw - nw - 0.2, shh - 0.1, [dict(t=l, size=9, color=TEXT, line=1.0)], anchor=MSO_ANCHOR.MIDDLE)

HX, HY, HD = 6.667, 4.55, 2.15
impacts = [("users_w", "Early identification", "of personnel needing welfare support — timely counselling & leave"),
           ("heart_w", "Fewer incidents", "reduction in stress-related incidents, self-harm and fratricide"),
           ("balance_w", "Balanced workload", "data-driven duty rosters, leave planning & transfer decisions"),
           ("retain_w", "Higher retention", "lower attrition / VRS; better job satisfaction & morale"),
           ("chartline_w", "Evidence-based welfare", "unit-level trends guide budgets, counsellor deployment & policy")]
benefits = [("smile_w", "Social", "confidential, destigmatised mental-health support for jawans & families"),
            ("medal_w", "Organisational", "enhanced force readiness, resilience & operational effectiveness"),
            ("coins_w", "Economic", "saves recruitment & training cost of every trained personnel retained"),
            ("flagin_w", "Strategic", "indigenous, sovereign welfare-tech for Indian forces (Atmanirbhar Bharat)"),
            ("expand_w", "Scalable", "State Police, NDRF / disaster response, government & corporate HR wellness")]
ry0, rh, rgap = 2.46, 0.78, 0.09
text(s, 0.4, 2.2, 4.9, 0.25, [dict(t="IMPACT", size=11, bold=True, color=SAFF, align=PP_ALIGN.CENTER)])
text(s, 8.05, 2.2, 4.9, 0.25, [dict(t="BENEFITS", size=11, bold=True, color=TEAL, align=PP_ALIGN.CENTER)])
for i in range(5):
    cy = ry0 + i*(rh + rgap) + rh/2
    line(s, 5.15, cy, HX, HY, "B8C1CC", 1.0, dash=MSO_LINE_DASH_STYLE.DASH)
    line(s, 8.18, cy, HX, HY, "B8C1CC", 1.0, dash=MSO_LINE_DASH_STYLE.DASH)
for i, (ic, t1, t2) in enumerate(impacts):
    y = ry0 + i*(rh + rgap)
    rrect(s, 0.4, y, 4.45, rh, SAFF_LT, radius=0.5)
    icon_circle(s, ic, 4.85 - 0.02, y + rh/2, 0.6, SAFF, 0.5)
    text(s, 0.55, y + 0.04, 3.85, rh - 0.08, [dict(runs=[dict(t=t1 + " — ", bold=True, color=NAVY, size=10.5), dict(t=t2, size=9.5, color=TEXT)], line=1.0, align=PP_ALIGN.RIGHT)], anchor=MSO_ANCHOR.MIDDLE)
for i, (ic, t1, t2) in enumerate(benefits):
    y = ry0 + i*(rh + rgap)
    rrect(s, 8.48, y, 4.45, rh, TEAL_LT, radius=0.5)
    icon_circle(s, ic, 8.48 + 0.02, y + rh/2, 0.6, TEAL, 0.5)
    text(s, 8.93, y + 0.04, 3.85, rh - 0.08, [dict(runs=[dict(t=t1 + " — ", bold=True, color=NAVY, size=10.5), dict(t=t2, size=9.5, color=TEXT)], line=1.0)], anchor=MSO_ANCHOR.MIDDLE)
oval(s, HX - HD/2 - 0.12, HY - HD/2 - 0.12, HD + 0.24, WHITE, LINE)
oval(s, HX - HD/2, HY - HD/2, HD, NAVY)
icon(s, "brain_big", HX - 0.3, HY - 0.85, 0.6)
text(s, HX - HD/2, HY - 0.2, HD, 1.1, [dict(t="SAHARA", size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER, after=2),
     dict(t="Reactive → Proactive\nWelfare Management", size=9, color="C9D3E6", align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.TOP)

# ---------------------------------------------------------------- SLIDE 6
s = S[5]
set_title(s, "RESEARCH  AND REFERENCES"); set_oval(s, TEAM_NAME); drop_shape(s, "TextBox 8")
groups = [
    ("gov_n", "Policy & Government Data", [
        ("MHA reply, Rajya Sabha (4 Dec 2024): 730 suicides, 47,891 VRS & 7,664 resignations in CAPFs/NSG/AR, 2020–24", "https://aninews.in/news/national/general-news/730-suicides-amongst-capfs-nsg-and-assam-rifles-during-the-last-five-years-nityanand-rai20241204183618/"),
        ("MHA Task Force on suicides & fratricide in CAPFs (2022) — service, working & personal stress factors", "https://www.deccanherald.com/india/mha-creates-taskforce-to-prevent-suicides-by-central-armed-police-forces-troops-1056571.html"),
        ("Digital Personal Data Protection Act, 2023 — MeitY (consent, purpose limitation, data-fiduciary duties)", "https://www.meity.gov.in/data-protection-framework"),
    ]),
    ("clinic_n", "Clinical Instruments & Existing Welfare Measures", [
        ("Kroenke et al. (2001) PHQ-9 depression scale; Spitzer et al. (2006) GAD-7 anxiety scale — validated screening tools", "https://pubmed.ncbi.nlm.nih.gov/11556941/"),
        ("Cohen et al. (1983) Perceived Stress Scale (PSS-10); WHO ICD-11 burn-out definition (2019)", "https://www.who.int/news/item/28-05-2019-burn-out-an-occupational-phenomenon-international-classification-of-diseases"),
        ("Tele-MANAS (MoHFW) national tele-mental-health helpline 14416 — counselling referral pathway; CRPF ‘Chaupal’ & buddy system", "https://telemanas.mohfw.gov.in/"),
    ]),
    ("book_n", "Research", [
        ("Schmidt et al. (2018) WESAD — wearable stress & affect detection dataset (ICMI)", "https://dl.acm.org/doi/10.1145/3242969.3242985"),
        ("Lundberg et al. (2020) TreeSHAP — exact local explanations for tree ensembles (Nature Machine Intelligence)", "https://www.nature.com/articles/s42256-019-0138-9"),
        ("Friedman (2001) Greedy function approximation: a gradient boosting machine (Annals of Statistics)", "https://projecteuclid.org/euclid.aos/1013203451"),
    ]),
    ("code_n2", "Technology & Standards", [
        ("Pedregosa et al. (2011) scikit-learn: Machine Learning in Python (JMLR) — training stack", "https://jmlr.org/papers/v12/pedregosa11a.html"),
        ("NIC MeghRaj Government Cloud & CERT-In security guidelines for hosting sensitive data", "https://cloud.gov.in/"),
        ("Android Health Connect API — consent-based access to wearable HRV, sleep & activity data", "https://developer.android.com/health-and-fitness/guides/health-connect"),
    ]),
]
gx = [0.4, 6.85]; gw = 6.1
gy = [1.3, 4.1]
for gi, (ic, gtitle, refs) in enumerate(groups):
    c, r = divmod(gi, 2)
    x = gx[c]; y = gy[r]
    rrect(s, x - 0.05, y - 0.08, gw + 0.1, 2.68, LIGHT, radius=0.04)
    icon_circle(s, ic, x + 0.2, y + 0.18, 0.36, WHITE, 0.5)
    text(s, x + 0.48, y, gw - 0.5, 0.36, [dict(t=gtitle, size=12.5, bold=True, color=NAVY)], anchor=MSO_ANCHOR.MIDDLE)
    ry = y + 0.45
    for j, (t, url) in enumerate(refs):
        text(s, x + 0.1, ry, gw - 0.15, 0.7, [
            dict(runs=[dict(t=f"{gi*3 + j + 1}. ", bold=True, color=SAFF, size=9.5), dict(t=t, size=9.5, bold=True, color=TEXT)], line=1.0, after=1),
            dict(runs=[dict(t=url, size=8, color=TEAL, url=url)], indent=0.2),
        ])
        ry += 0.78

prs.save(OUT)
print("saved", OUT)
